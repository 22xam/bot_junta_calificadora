import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os

# Assuming extractor.py and main.py are in the same directory or PYTHONPATH
from extractor import extraer_por_codigo
from main import save_extracted_data

# Imports for DOCX generation
from docx import Document
from llenar_docx import cargar_json, detectar_campos, reemplazar_campos, completar_listado


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Extractor de Datos PDF y Generador DOCX") # Translated

        # PDF Processing related variables
        self.pdf_path = tk.StringVar()
        self.subject_code_var = tk.StringVar()
        self.current_json_path = None # Stores path to the JSON from extraction step

        # DOCX Generation related variables
        self.docx_template_path = tk.StringVar()
        self.docx_dynamic_fields_entries = {} # Stores Entry StringVars for template fields

        # General status
        self.status_var = tk.StringVar()

        self.root.geometry("800x850") # Adjusted height slightly

        # Main frame
        main_frame = ttk.Frame(self.root, padding="10 10 10 10") # Uniform padding
        main_frame.pack(fill=tk.BOTH, expand=True)

        # PDF Processing Frame
        pdf_frame = ttk.LabelFrame(main_frame, text="1. Procesamiento PDF", padding="10 10 10 10") # Translated
        pdf_frame.pack(fill=tk.X, pady=10)

        # PDF File Selection
        select_pdf_button = ttk.Button(pdf_frame, text="Seleccionar PDF", command=self.select_pdf_file) # Translated
        select_pdf_button.grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.pdf_path_label = ttk.Label(pdf_frame, text="Ningún PDF seleccionado") # Translated
        self.pdf_path_label.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        # Subject Code Input
        subject_code_label = ttk.Label(pdf_frame, text="Código de Materia:") # Translated
        subject_code_label.grid(row=1, column=0, padx=5, pady=5, sticky="w")
        subject_code_entry = ttk.Entry(pdf_frame, textvariable=self.subject_code_var, width=30)
        subject_code_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")

        # Extract Data Button
        extract_button = ttk.Button(pdf_frame, text="Extraer Datos del PDF", command=self.process_extraction) # Translated
        extract_button.grid(row=2, column=0, columnspan=2, pady=(10,5))

        pdf_frame.grid_columnconfigure(1, weight=1)

        # Results Display Frame (for Treeview)
        results_frame = ttk.LabelFrame(main_frame, text="2. Visor de Datos Extraídos", padding="10 10 10 10") # Translated
        results_frame.pack(fill=tk.BOTH, expand=True, pady=10)

        self.results_treeview = ttk.Treeview(results_frame, columns=("Orden", "Apellido y Nombre", "Puntos"), show="headings")
        self.results_treeview.heading("Orden", text="Orden") # Already Spanish or neutral
        self.results_treeview.heading("Apellido y Nombre", text="Apellido y Nombre") # Already Spanish
        self.results_treeview.heading("Puntos", text="Puntos") # Already Spanish

        self.results_treeview.column("Orden", width=60, anchor=tk.CENTER, stretch=tk.NO)
        self.results_treeview.column("Apellido y Nombre", width=350, stretch=tk.YES)
        self.results_treeview.column("Puntos", width=100, anchor=tk.CENTER, stretch=tk.NO)

        # Scrollbars for Treeview
        vsb = ttk.Scrollbar(results_frame, orient="vertical", command=self.results_treeview.yview)
        hsb = ttk.Scrollbar(results_frame, orient="horizontal", command=self.results_treeview.xview)
        self.results_treeview.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        # Use grid for treeview and scrollbars to manage space better
        results_frame.grid_rowconfigure(0, weight=1)
        results_frame.grid_columnconfigure(0, weight=1)

        self.results_treeview.grid(row=0, column=0, sticky="nsew", padx=(0,5), pady=(0,5)) # Pad to not touch scrollbars
        vsb.grid(row=0, column=1, sticky="ns", pady=(0,5))
        hsb.grid(row=1, column=0, sticky="ew", padx=(0,5))


        # DOCX Generation Frame
        docx_frame = ttk.LabelFrame(main_frame, text="3. Generación DOCX", padding="10 10 10 10") # Translated
        docx_frame.pack(fill=tk.X, pady=10, expand=False) # Changed expand to False, fixed height for canvas container

        # Template Selection
        select_template_button = ttk.Button(docx_frame, text="Seleccionar Plantilla DOCX", command=self.select_docx_template) # Translated
        select_template_button.grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.docx_template_path_label = ttk.Label(docx_frame, text="Ninguna plantilla DOCX seleccionada") # Translated
        self.docx_template_path_label.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        # --- Scrollable Frame for Dynamic Fields ---
        # Container for Canvas and Scrollbar
        dynamic_fields_container = ttk.Frame(docx_frame)
        dynamic_fields_container.grid(row=1, column=0, columnspan=2, sticky="nsew", pady=10, padx=5)
        dynamic_fields_container.grid_rowconfigure(0, weight=1)
        dynamic_fields_container.grid_columnconfigure(0, weight=1)
        # Set a fixed height for the scrollable area, or allow it to expand
        # For fixed height:
        dynamic_fields_container.config(height=200) # Example fixed height

        self.fields_canvas = tk.Canvas(dynamic_fields_container)
        self.fields_canvas.grid(row=0, column=0, sticky="nsew")

        self.fields_scrollbar = ttk.Scrollbar(dynamic_fields_container, orient="vertical", command=self.fields_canvas.yview)
        self.fields_scrollbar.grid(row=0, column=1, sticky="ns")

        self.fields_canvas.configure(yscrollcommand=self.fields_scrollbar.set)

        # This is the frame that will contain the actual labels and entries
        self.dynamic_fields_frame = ttk.LabelFrame(self.fields_canvas, text="Campos de Plantilla", padding="10 10 10 10") # Translated
        self.dynamic_fields_window_id = self.fields_canvas.create_window((0, 0), window=self.dynamic_fields_frame, anchor="nw")

        # Bindings to update scrollregion and frame width
        self.dynamic_fields_frame.bind("<Configure>", self._on_dynamic_fields_frame_configure)
        self.fields_canvas.bind("<Configure>", self._on_fields_canvas_configure)
        # --- End Scrollable Frame ---

        # Generate Document Button
        self.generate_doc_button = ttk.Button(docx_frame, text="Generar Documento desde Plantilla", command=self.generate_document) # Translated
        self.generate_doc_button.grid(row=2, column=0, columnspan=2, pady=(10,5))
        self.generate_doc_button.config(state=tk.DISABLED)

        docx_frame.grid_columnconfigure(1, weight=1)
        docx_frame.grid_rowconfigure(1, weight=1) # Allow the dynamic_fields_container to expand vertically if needed (if not fixed height)


        # Status Frame / Bar
        status_frame = ttk.LabelFrame(main_frame, text="Estado de la Aplicación", padding="10 10 10 10") # Translated
        status_frame.pack(fill=tk.X, pady=(10,0))
        self.status_label = ttk.Label(status_frame, textvariable=self.status_var, wraplength=780)
        self.status_label.pack(fill=tk.X, padx=5, pady=5)
        self.status_var.set("Listo. Seleccione PDF e ingrese Código de Materia para iniciar.") # Translated

    def _on_dynamic_fields_frame_configure(self, event=None):
        self.fields_canvas.configure(scrollregion=self.fields_canvas.bbox("all"))

    def _on_fields_canvas_configure(self, event=None):
        self.fields_canvas.itemconfig(self.dynamic_fields_window_id, width=event.width)

    def select_pdf_file(self):
        filepath = filedialog.askopenfilename(
            title="Seleccionar Archivo PDF", # Translated
            filetypes=(("Archivos PDF", "*.pdf"), ("Todos los archivos", "*.*")) # Translated
        )
        if filepath:
            self.pdf_path.set(filepath)
            self.pdf_path_label.config(text=os.path.basename(filepath))
            self.status_var.set(f"PDF seleccionado: {os.path.basename(filepath)}") # Translated
        else:
            self.pdf_path.set("")
            self.pdf_path_label.config(text="Ningún PDF seleccionado") # Translated
            self.status_var.set("Selección de PDF cancelada.") # Translated
        self.check_enable_generate_button()

    def process_extraction(self):
        self.status_var.set("Procesando...") # Translated
        self.current_json_path = None # Reset
        self.check_enable_generate_button() # Update button state

        # Clear previous results from Treeview
        for item in self.results_treeview.get_children():
            self.results_treeview.delete(item)

        pdf_path_str = self.pdf_path.get()
        subject_code_str = self.subject_code_var.get().strip()

        if not pdf_path_str or not os.path.exists(pdf_path_str):
            self.status_var.set("Error: Por favor, seleccione un archivo PDF válido.") # Translated
            messagebox.showerror("Error de Entrada", "Por favor, seleccione un archivo PDF válido.") # Translated
            return

        if not subject_code_str:
            self.status_var.set("Error: Se requiere el Código de Materia.") # Translated
            messagebox.showerror("Error de Entrada", "Se requiere el Código de Materia.") # Translated
            return

        try:
            self.status_var.set(f"Extrayendo datos para el código '{subject_code_str}' de '{os.path.basename(pdf_path_str)}'...") # Translated
            self.root.update_idletasks()

            nombre_materia, registros = extraer_por_codigo(pdf_path_str, subject_code_str)

            if nombre_materia is None or not nombre_materia.strip():
                self.status_var.set(f"Código de materia '{subject_code_str}' no encontrado en el PDF.") # Translated
                messagebox.showinfo("Información de Extracción", f"Código de materia '{subject_code_str}' no encontrado en el PDF.") # Translated
                return

            if not registros:
                self.status_var.set(f"No se encontraron registros de alumnos para la materia '{nombre_materia}' (Código: {subject_code_str}).") # Translated
                messagebox.showinfo("Información de Extracción", f"No se encontraron registros de alumnos para la materia '{nombre_materia}' (Código: {subject_code_str}).") # Translated
                return

            for reg in registros:
                self.results_treeview.insert("", tk.END, values=(reg.get("Orden"), reg.get("Apellido y Nombre"), reg.get("Puntos")))

            self.status_var.set(f"Se encontraron {len(registros)} registros para '{nombre_materia}'. Guardando datos...") # Translated
            self.root.update_idletasks()

            output_directory = "output_data"
            os.makedirs(output_directory, exist_ok=True)
            csv_path, json_path = save_extracted_data(subject_code_str, nombre_materia, registros, output_dir=output_directory)
            self.current_json_path = json_path

            self.status_var.set(f"{len(registros)} registros extraídos. Guardado en carpeta '{output_directory}'. JSON: {os.path.basename(json_path)}") # Translated
            messagebox.showinfo("Extracción Exitosa", f"{len(registros)} registros extraídos para '{nombre_materia}'.\nArchivos guardados en carpeta '{output_directory}'.\nCSV: {os.path.basename(csv_path)}\nJSON: {os.path.basename(json_path)}") # Translated

        except Exception as e:
            self.status_var.set(f"Error durante la extracción del PDF: {e}") # Translated (clarified source)
            messagebox.showerror("Error de Extracción PDF", f"Ocurrió un error durante la extracción de datos del PDF:\n{e}") # Translated
        finally:
            self.check_enable_generate_button()

    def select_docx_template(self):
        filepath = filedialog.askopenfilename(
            title="Seleccionar Archivo de Plantilla DOCX", # Translated
            filetypes=(("Documentos Word", "*.docx"), ("Todos los archivos", "*.*")) # Translated
        )
        if filepath:
            self.docx_template_path.set(filepath)
            self.docx_template_path_label.config(text=os.path.basename(filepath))
            self.status_var.set(f"Plantilla DOCX seleccionada: {os.path.basename(filepath)}") # Translated
            self.populate_dynamic_fields(filepath)
        else:
            self.docx_template_path.set("")
            self.docx_template_path_label.config(text="Ninguna plantilla DOCX seleccionada") # Translated
            for widget in self.dynamic_fields_frame.winfo_children():
                widget.destroy()
            self.docx_dynamic_fields_entries.clear()
            self.status_var.set("Selección de plantilla DOCX cancelada.") # Translated
        self.check_enable_generate_button()

    def populate_dynamic_fields(self, template_path):
        for widget in self.dynamic_fields_frame.winfo_children():
            widget.destroy()
        self.docx_dynamic_fields_entries.clear()

        try:
            doc = Document(template_path)
            campos = detectar_campos(doc)

            if not campos:
                no_fields_label = ttk.Label(self.dynamic_fields_frame, text="No se encontraron campos personalizados (ej: _CAMPO) en la plantilla.") # Translated
                no_fields_label.pack(pady=10, padx=5)
                return

            title_label = ttk.Label(self.dynamic_fields_frame, text="Complete los campos de la plantilla:", font="-weight bold") # Translated
            title_label.pack(pady=(5,10), padx=5, anchor="w")

            for idx, field_name in enumerate(sorted(list(campos))):
                if field_name.upper() == "_MATERIA":
                    continue

                field_row_frame = ttk.Frame(self.dynamic_fields_frame)
                field_row_frame.pack(fill=tk.X, pady=3, padx=5)

                display_name = field_name.replace("_", " ").strip().title()

                label = ttk.Label(field_row_frame, text=f"{display_name}:", width=25)
                label.pack(side=tk.LEFT, padx=(0,5))

                entry_var = tk.StringVar()
                entry = ttk.Entry(field_row_frame, textvariable=entry_var)
                entry.pack(side=tk.LEFT, expand=True, fill=tk.X)

                self.docx_dynamic_fields_entries[field_name] = entry_var

        except Exception as e:
            self.status_var.set(f"Error procesando plantilla DOCX: {e}") # Translated
            messagebox.showerror("Error de Plantilla", f"Error procesando plantilla DOCX. Asegúrese que es un archivo .docx válido.\nDetalle: {e}") # Translated
            # No need to loop destroy here, new populate_dynamic_fields starts by clearing.
            self.docx_dynamic_fields_entries.clear() # ensure entries are cleared
        # Manually trigger a configure event for the canvas if fields were populated
        # to ensure scrollregion is set correctly, especially if window hasn't been resized yet.
        self.dynamic_fields_frame.update_idletasks() # Ensure frame has its new size
        self._on_dynamic_fields_frame_configure()


    def generate_document(self):
        if not self.current_json_path or not os.path.exists(self.current_json_path):
            self.status_var.set("Error: Datos extraídos (JSON) no encontrados. Por favor, extraiga datos primero.") # Translated
            messagebox.showerror("Datos Faltantes", "Datos extraídos (JSON) no encontrados. Por favor, ejecute 'Extraer Datos del PDF' primero.") # Translated
            return

        template_path_str = self.docx_template_path.get()
        if not template_path_str or not os.path.exists(template_path_str):
            self.status_var.set("Error: Plantilla DOCX no seleccionada o no encontrada.") # Translated
            messagebox.showerror("Plantilla Faltante", "Plantilla DOCX no seleccionada o no encontrada. Por favor, seleccione una plantilla.") # Translated
            return

        self.status_var.set("Generando documento...") # Translated
        self.root.update_idletasks()

        try:
            json_data = cargar_json(self.current_json_path)

            user_provided_data = {}
            for field_name, entry_var in self.docx_dynamic_fields_entries.items():
                user_provided_data[field_name] = entry_var.get()

            final_datos_for_docx = {
                "_MATERIA": json_data.get("materia", "MATERIA NO ESPECIFICADA") # Keep this default or make it Spanish
            }
            final_datos_for_docx.update(user_provided_data)

            temp_doc_check = Document(template_path_str)
            all_template_fields = detectar_campos(temp_doc_check)
            for field in all_template_fields:
                if field not in final_datos_for_docx:
                    final_datos_for_docx[field] = ""

            doc = Document(template_path_str)
            reemplazar_campos(doc, final_datos_for_docx)
            completar_listado(doc, json_data.get("listado", []))

            output_docs_dir = "output_documents"
            os.makedirs(output_docs_dir, exist_ok=True)

            base_name = os.path.basename(self.current_json_path)
            name_part = base_name.replace("_listado.json", "").replace("_resumido.csv","")
            output_filename = f"{name_part}_designacion.docx"
            output_docx_path = os.path.join(output_docs_dir, output_filename)

            doc.save(output_docx_path)
            self.status_var.set(f"Documento generado con éxito: {output_docx_path}") # Translated
            messagebox.showinfo("Éxito", f"¡Documento generado con éxito!\nGuardado en carpeta '{output_docs_dir}' como: {output_filename}") # Translated

        except FileNotFoundError as e:
            self.status_var.set(f"Error: Archivo no encontrado. {e}") # Translated
            messagebox.showerror("Error de Archivo No Encontrado", f"Un archivo requerido no fue encontrado.\n{e}") # Translated
        except ValueError as e:
            self.status_var.set(f"Error: Formato de datos inválido. {e}") # Translated
            messagebox.showerror("Error de Formato de Datos", f"Hubo un problema con el formato de los datos.\n{e}") # Translated
        except IOError as e:
            self.status_var.set(f"Error: Falló la operación de archivo. {e}") # Translated
            messagebox.showerror("Error de Operación de Archivo", f"Falló una operación de archivo.\n{e}") # Translated
        except Exception as e:
            self.status_var.set(f"Error generando documento: {e}") # Translated
            messagebox.showerror("Error de Generación", f"Ocurrió un error mientras se generaba el documento:\n{e}") # Translated
        finally:
            self.check_enable_generate_button()

    def check_enable_generate_button(self):
        if self.current_json_path and \
           os.path.exists(self.current_json_path) and \
           self.docx_template_path.get() and \
           os.path.exists(self.docx_template_path.get()):
            self.generate_doc_button.config(state=tk.NORMAL)
        else:
            self.generate_doc_button.config(state=tk.DISABLED)


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
